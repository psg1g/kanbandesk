from datetime import date, datetime

from django.http import HttpResponse, JsonResponse
from django.views.decorators.csrf import csrf_exempt
from rest_framework import viewsets
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import AllowAny, IsAuthenticated
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MAX_EXPORT_ROWS = 10_000


def _safe_cell(value):
    """Защита от инъекций формул: значения =, +, -, @ пишутся с апострофом."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


class ExcelExportMixin:
    """GET /api/boards/export/ — выгрузка записей пользователя в .xlsx.

    Лист «Доски» — сводка по доскам; лист «Карточки» — подробные данные
    каждой карточки (колонка, приоритет, срок, вложения и т.д.).
    """

    export_columns = []  # list of (header, getter)
    detail_columns = []  # list of (header, getter) для листа «Карточки»

    @action(detail=False, methods=["get"], url_path="export")
    def export(self, request):
        qs = self.filter_queryset(self.get_queryset())
        count = qs.count()
        truncated = count > MAX_EXPORT_ROWS
        rows = qs[:MAX_EXPORT_ROWS]

        wb = Workbook()
        ws = wb.active
        ws.title = "Доски"
        headers = [h for h, _ in self.export_columns]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
        )

        for obj in rows:
            r = ws.max_row + 1
            for col_idx, (_, getter) in enumerate(self.export_columns, start=1):
                v = getter(obj)
                cell = ws.cell(row=r, column=col_idx)
                if isinstance(v, (date, datetime)):
                    cell.value = v
                    cell.number_format = "DD.MM.YYYY"
                elif isinstance(v, (int, float)):
                    cell.value = v
                else:
                    cell.value = _safe_cell(v)

        # ширина колонок под содержимое
        for idx, header in enumerate(headers, start=1):
            letter = get_column_letter(idx)
            width = len(str(header))
            for r in range(2, ws.max_row + 1):
                val = ws.cell(row=r, column=idx).value
                if val is not None:
                    width = max(width, min(len(str(val)), 60))
            ws.column_dimensions[letter].width = width + 2

        # --- Лист «Карточки»: подробные данные по всем карточкам досок ---
        if self.detail_columns:
            ws2 = wb.create_sheet("Карточки")
            d_headers = [h for h, _ in self.detail_columns]
            ws2.append(d_headers)
            for cell in ws2[1]:
                cell.font = Font(bold=True)
            ws2.freeze_panes = "A2"

            detail_count = 0
            for board in rows:
                for col in board.columns.all():
                    for card in col.cards.all():
                        if detail_count >= MAX_EXPORT_ROWS:
                            break
                        r = ws2.max_row + 1
                        for col_idx, (_, getter) in enumerate(
                            self.detail_columns, start=1
                        ):
                            v = getter(board, col, card)
                            cell = ws2.cell(row=r, column=col_idx)
                            if isinstance(v, (date, datetime)):
                                cell.value = v
                                cell.number_format = "DD.MM.YYYY"
                            elif isinstance(v, (int, float)):
                                cell.value = v
                            else:
                                cell.value = _safe_cell(v)
                        detail_count += 1
            ws2.auto_filter.ref = (
                f"A1:{get_column_letter(len(d_headers))}{max(ws2.max_row, 1)}"
            )
            for idx, header in enumerate(d_headers, start=1):
                letter = get_column_letter(idx)
                width = len(str(header))
                for r in range(2, ws2.max_row + 1):
                    val = ws2.cell(row=r, column=idx).value
                    if val is not None:
                        width = max(width, min(len(str(val)), 60))
                ws2.column_dimensions[letter].width = width + 2

        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        filename = f"kanban_{date.today():%Y-%m-%d}.xlsx"
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        if truncated:
            response["X-Export-Warning"] = (
                f"Выгружено только первые {MAX_EXPORT_ROWS} из {count} записей"
            )
        wb.save(response)
        return response
